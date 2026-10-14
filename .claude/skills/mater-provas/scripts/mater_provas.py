#!/usr/bin/env python3
"""Testes e provas no modelo oficial da Escola Mater / Colégio São Conrado.

O gerador não recria o layout: abre o próprio MODELO AVALIAÇÕES (Drive), preenche
somente os campos em vermelho do cabeçalho, remove o conteúdo de exemplo e escreve
as questões com a formatação do corpo do modelo (Tahoma 11, entrelinha 1,15).

Subcomandos:
  modelo  decodifica o modelo baixado pelo MCP do Google Drive e confere a geometria oficial
  gerar   monta DOCX + PDF do aluno, gabarito do professor, prévias PNG e manifesto, com gate de entrega
  figura  renderiza uma ilustração SVG em PNG de alta resolução, em tons de cinza para impressão P&B
"""
import argparse
import base64
import copy
import datetime as dt
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from docx.text.run import Run
from lxml import etree

DRIVE_FILE_ID = "1ZfmTSo7m8h9kpnhzaut7GaoN6jAor9sY"
DRIVE_TITULO = "MODELO AVALIAÇÕES - Fund II e Médio.docx"

PAGINA = {"w": "11906", "h": "16838"}
MARGENS = {"top": "426", "right": "707", "bottom": "720", "left": "709"}
GRADE_CABECALHO = ["4943", "851", "2126", "2838"]

# texto do campo em vermelho no modelo -> chave preenchida a partir da especificação
CAMPOS = ["XºAno", "/Xª Série", "MATÉRIA", "  /03/2026", "SEU NOME", "TESTE/PROVA X", "3,0/7,0"]

FONTE_CORPO = "Tahoma"
TAMANHO_CORPO = 22  # meios-pontos (11 pt)
RECUO_ALTERNATIVA = 284  # twips, igual às alternativas do modelo
ENTRELINHA = 1.15  # w:line=276 lineRule=auto

LETRAS = "abcdefgh"
SEGMENTOS_COM_SITUACAO_PROBLEMA = {"fundamental2", "fundamental_ii", "medio", "ensino_medio"}


class ErroModelo(Exception):
    pass


# ---------------------------------------------------------------- utilidades

def sha256(caminho):
    return hashlib.sha256(Path(caminho).read_bytes()).hexdigest()


def decimal_br(valor, campo):
    try:
        return Decimal(str(valor).strip().replace(",", "."))
    except InvalidOperation:
        raise ValueError(f"{campo}: valor numérico inválido {valor!r}")


def formatar_br(d):
    texto = format(d.normalize(), "f")
    if "." not in texto:
        texto += ".0"
    return texto.replace(".", ",")


def titulo_disciplina(texto):
    return " ".join(p if p.isdigit() else p.capitalize() for p in texto.split())


def nome_base(spec):
    # regra do modelo: "Turma(Série/Ano) – Avaliação – Disciplina"
    return spec.get("nome_arquivo") or f"{spec['serie']} – Avaliação – {titulo_disciplina(spec['disciplina'])}"


def paragrafos_de(texto):
    if isinstance(texto, list):
        return [str(t).strip() for t in texto if str(t).strip()]
    return [t.strip() for t in re.split(r"\n\s*\n", str(texto or "")) if t.strip()]


def carregar_spec(caminho):
    texto = Path(caminho).read_text(encoding="utf-8")
    if Path(caminho).suffix.lower() in {".yaml", ".yml"}:
        import yaml
        return yaml.safe_load(texto)
    return json.loads(texto)


# ---------------------------------------------------------------- modelo oficial

def decodificar_modelo(origem, destino):
    """Aceita o .docx ou o arquivo JSON salvo pelo download_file_content do MCP do Drive."""
    dados = Path(origem).read_bytes()
    if not dados.startswith(b"PK"):
        conteudo = json.loads(dados.decode("utf-8"))
        if isinstance(conteudo, dict) and "content" in conteudo:
            conteudo = conteudo["content"]
        dados = base64.b64decode(conteudo)
    if not dados.startswith(b"PK"):
        raise ErroModelo("o arquivo decodificado não é um .docx")
    Path(destino).parent.mkdir(parents=True, exist_ok=True)
    Path(destino).write_bytes(dados)


def conferir_modelo(caminho):
    erros, avisos = [], []
    doc = Document(caminho)
    sect = doc.sections[0]._sectPr
    pg = sect.find(qn("w:pgSz"))
    if pg is None or pg.get(qn("w:w")) != PAGINA["w"] or pg.get(qn("w:h")) != PAGINA["h"]:
        avisos.append("tamanho de página diferente do A4 do modelo conhecido")
    mar = sect.find(qn("w:pgMar"))
    if mar is None or any(mar.get(qn(f"w:{k}")) != v for k, v in MARGENS.items()):
        avisos.append("margens diferentes das do modelo conhecido (0,75/1,25/1,27/1,25 cm)")
    if sect.find(qn("w:pgBorders")) is None:
        avisos.append("modelo sem borda de página")
    if not doc.tables:
        erros.append("tabela do cabeçalho não encontrada")
        return erros, avisos
    grade = [g.get(qn("w:w")) for g in doc.tables[0]._tbl.tblGrid.findall(qn("w:gridCol"))]
    if grade != GRADE_CABECALHO:
        avisos.append(f"grade da tabela do cabeçalho mudou: {grade}")
    contagem = {c: 0 for c in CAMPOS}
    for run in runs_da_tabela(doc.tables[0]):
        if run.text in contagem:
            contagem[run.text] += 1
    for campo, n in contagem.items():
        if n != 1:
            erros.append(f"campo do cabeçalho {campo!r} aparece {n} vez(es) (esperado 1)")
    textos = [p.text.strip() for p in doc.paragraphs]
    for titulo in ("ESCOLA MATER", "COLÉGIO SÃO CONRADO"):
        if titulo not in textos:
            erros.append(f"título {titulo!r} não encontrado")
    if not any(t.startswith("Questão 1") for t in textos):
        avisos.append("questão de exemplo não encontrada; não foi possível conferir a fonte do corpo")
    else:
        p = next(p for p in doc.paragraphs if p.text.strip().startswith("Questão 1"))
        fontes = {r._r.rPr.rFonts.get(qn("w:ascii")) for r in p.runs if r._r.rPr is not None and r._r.rPr.rFonts is not None}
        if fontes != {FONTE_CORPO}:
            avisos.append(f"fonte do corpo no modelo agora é {fontes}; o gerador usa {FONTE_CORPO}")
    if cabecalho_com_logo(doc) is None:
        erros.append("cabeçalho de página com o logo não encontrado")
    return erros, avisos


def runs_da_tabela(tabela):
    # percorre o XML direto: row.cells repete células mescladas (gridSpan)
    for r in tabela._tbl.iter(qn("w:r")):
        yield Run(r, tabela)


def cabecalho_com_logo(doc):
    rels = doc.part.rels
    for ref in doc.sections[0]._sectPr.findall(qn("w:headerReference")):
        parte = rels[ref.get(qn("r:id"))].target_part
        if parte.element.findall(".//" + qn("w:drawing")):
            return ref.get(qn("r:id"))
    return None


# ---------------------------------------------------------------- cabeçalho

def preencher_cabecalho(doc, spec, preto):
    tipo = str(spec.get("tipo", "TESTE")).upper()
    valores = {
        "XºAno": spec["serie"],
        "/Xª Série": "",
        "MATÉRIA": spec["disciplina"].upper(),
        "  /03/2026": spec.get("data") or f"___/___/{spec.get('ano', dt.date.today().year)}",
        "SEU NOME": spec["professor"],
        "TESTE/PROVA X": f"{tipo} {spec['bimestre']}",
        "3,0/7,0": formatar_br(decimal_br(spec["total_pontos"], "total_pontos")),
    }
    usados = set()
    for run in runs_da_tabela(doc.tables[0]):
        if run.text in valores and run.text not in usados:
            usados.add(run.text)
            run.text = valores[run.text]
            if preto:
                run.font.color.rgb = RGBColor(0, 0, 0)
    faltando = set(valores) - usados
    if faltando:
        raise ErroModelo(f"campos do cabeçalho não encontrados no modelo: {sorted(faltando)}")


def cabecalho_so_na_primeira_pagina(doc):
    """O logo do cabeçalho de página fica só na 1ª página (no modelo ele se repete e invade o texto)."""
    sect = doc.sections[0]._sectPr
    rid_logo = cabecalho_com_logo(doc)
    primeira = None
    for ref in sect.findall(qn("w:headerReference")):
        if ref.get(qn("w:type")) == "first":
            primeira = ref
        elif ref.get(qn("w:type")) == "default":
            sect.remove(ref)
    if primeira is None:
        primeira = OxmlElement("w:headerReference")
        primeira.set(qn("w:type"), "first")
        sect.insert(0, primeira)
    primeira.set(qn("r:id"), rid_logo)
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.header.is_linked_to_previous = False  # cria um cabeçalho padrão vazio
    rodape = next((r for r in sect.findall(qn("w:footerReference")) if r.get(qn("w:type")) == "default"), None)
    if rodape is not None and not any(r.get(qn("w:type")) == "first" for r in sect.findall(qn("w:footerReference"))):
        primeiro = copy.deepcopy(rodape)
        primeiro.set(qn("w:type"), "first")
        rodape.addnext(primeiro)


def remover_exemplo_do_corpo(doc):
    corpo = doc.element.body
    filhos = list(corpo)
    alvo = next((i for i, el in enumerate(filhos)
                 if el.tag == qn("w:p") and "COLÉGIO SÃO CONRADO" in "".join(el.itertext())), None)
    if alvo is None:
        raise ErroModelo("parágrafo 'COLÉGIO SÃO CONRADO' não encontrado")
    for el in filhos[alvo + 1:]:
        if el.tag in (qn("w:p"), qn("w:tbl")):
            corpo.remove(el)


# ---------------------------------------------------------------- texto e fórmulas

MARCAS = re.compile(
    r"\\\((?P<math>.+?)\\\)"
    r"|\*\*(?P<b>.+?)\*\*"
    r"|\*(?P<i>[^*\n]+?)\*"
    r"|\^\{(?P<sup>[^}]*)\}"
    r"|_\{(?P<sub>[^}]*)\}"
)
ESCAPES = {"\\*": "\ue000", "\\^": "\ue001", "\\_": "\ue002"}


def restaurar_escapes(texto):
    for marca, substituto in ESCAPES.items():
        texto = texto.replace(substituto, marca)
    return texto


def formulas_do_texto(texto):
    return [m.group("math").strip() for m in MARCAS.finditer(texto or "") if m.group("math")]


def converter_formulas(latex):
    """Converte LaTeX em OMML (equação nativa do Word) numa única chamada ao pandoc."""
    latex = list(dict.fromkeys(latex))
    if not latex:
        return {}
    if not shutil.which("pandoc"):
        raise ErroModelo("há fórmulas \\( \\) na especificação, mas o pandoc não está instalado")
    markdown = "\n\n".join(f"${f}$" for f in latex)
    with tempfile.TemporaryDirectory() as tmp:
        saida = Path(tmp) / "formulas.docx"
        subprocess.run(["pandoc", "-f", "markdown", "-t", "docx", "-o", str(saida)],
                       input=markdown.encode("utf-8"), check=True, capture_output=True)
        xml = zipfile.ZipFile(saida).read("word/document.xml")
    paragrafos = etree.fromstring(xml).find(qn("w:body")).findall(qn("w:p"))
    resultado = {}
    for formula, p in zip(latex, paragrafos):
        omath = p.find(".//" + qn("m:oMath"))
        if omath is None:
            raise ErroModelo(f"o pandoc não conseguiu converter a fórmula: {formula}")
        for mr in omath.iter(qn("m:r")):
            rpr = mr.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                mrpr = mr.find(qn("m:rPr"))
                mr.insert(0 if mrpr is None else 1, rpr)
            for tag in ("w:sz", "w:szCs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = OxmlElement(tag)
                    rpr.append(el)
                el.set(qn("w:val"), str(TAMANHO_CORPO))
        resultado[formula] = omath
    return resultado


def propriedades_run(negrito=False, italico=False, vertical=None):
    rpr = OxmlElement("w:rPr")
    fontes = OxmlElement("w:rFonts")
    for atributo in ("ascii", "eastAsia", "hAnsi", "cs"):
        fontes.set(qn(f"w:{atributo}"), FONTE_CORPO)
    rpr.append(fontes)
    if negrito:
        rpr.append(OxmlElement("w:b"))
        rpr.append(OxmlElement("w:bCs"))
    if italico:
        rpr.append(OxmlElement("w:i"))
        rpr.append(OxmlElement("w:iCs"))
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), str(TAMANHO_CORPO))
        rpr.append(el)
    if vertical:
        el = OxmlElement("w:vertAlign")
        el.set(qn("w:val"), vertical)
        rpr.append(el)
    return rpr


def adicionar_run(p, texto, negrito=False, italico=False, vertical=None):
    for i, linha in enumerate(texto.split("\n")):
        if i:
            r = p.add_run()
            r._r.insert(0, propriedades_run(negrito, italico, vertical))
            r.add_break()
        if linha:
            for marca, original in ESCAPES.items():
                linha = linha.replace(original, marca[1])
            r = p.add_run(linha)
            r._r.insert(0, propriedades_run(negrito, italico, vertical))


def escrever_texto(p, texto, formulas, negrito=False, italico=False):
    for marca, substituto in ESCAPES.items():
        texto = texto.replace(marca, substituto)
    pos = 0
    for m in MARCAS.finditer(texto):
        if m.start() > pos:
            adicionar_run(p, texto[pos:m.start()], negrito, italico)
        if m.group("math"):
            p._p.append(copy.deepcopy(formulas[restaurar_escapes(m.group("math").strip())]))
        elif m.group("b"):
            escrever_texto(p, m.group("b"), formulas, True, italico)
        elif m.group("i"):
            escrever_texto(p, m.group("i"), formulas, negrito, True)
        elif m.group("sup") is not None:
            adicionar_run(p, m.group("sup"), negrito, italico, "superscript")
        else:
            adicionar_run(p, m.group("sub"), negrito, italico, "subscript")
        pos = m.end()
    if pos < len(texto):
        adicionar_run(p, texto[pos:], negrito, italico)


# ---------------------------------------------------------------- corpo

class Corpo:
    def __init__(self, doc, formulas, base):
        self.doc = doc
        self.formulas = formulas
        self.base = base
        numeracao = doc.part.numbering_part.element
        self.numeracao = numeracao
        self.abstrato_letras = next(
            (a.get(qn("w:abstractNumId")) for a in numeracao.findall(qn("w:abstractNum"))
             if a.find(qn("w:lvl")).find(qn("w:numFmt")).get(qn("w:val")) == "lowerLetter"),
            None)
        if self.abstrato_letras is None:
            raise ErroModelo("lista 'a)' do modelo não encontrada em numbering.xml")

    def nova_lista(self):
        """Cada questão reinicia as alternativas em a)."""
        ids = [int(n.get(qn("w:numId"))) for n in self.numeracao.findall(qn("w:num"))]
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(max(ids) + 1))
        abstrato = OxmlElement("w:abstractNumId")
        abstrato.set(qn("w:val"), self.abstrato_letras)
        num.append(abstrato)
        sobrescrita = OxmlElement("w:lvlOverride")
        sobrescrita.set(qn("w:ilvl"), "0")
        inicio = OxmlElement("w:startOverride")
        inicio.set(qn("w:val"), "1")
        sobrescrita.append(inicio)
        num.append(sobrescrita)
        limpeza = self.numeracao.find(qn("w:numIdMacAtCleanup"))
        if limpeza is None:
            self.numeracao.append(num)
        else:
            limpeza.addprevious(num)
        return max(ids) + 1

    def paragrafo(self, alinhamento=None, manter_com_proximo=False, lista=None):
        p = self.doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.line_spacing = ENTRELINHA
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        if manter_com_proximo:
            fmt.keep_with_next = True
            fmt.keep_together = True
        if lista is not None:
            num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = lista
            fmt.left_indent = Twips(RECUO_ALTERNATIVA)
            fmt.first_line_indent = Twips(-RECUO_ALTERNATIVA)
        if alinhamento is not None:
            fmt.alignment = alinhamento
        marca = propriedades_run()
        p._p.get_or_add_pPr().append(marca)
        return p

    def linhas_em_branco(self, n, manter_com_proximo=False):
        for _ in range(n):
            self.paragrafo(manter_com_proximo=manter_com_proximo)

    def imagem(self, img):
        p = self.paragrafo(WD_ALIGN_PARAGRAPH.CENTER, manter_com_proximo=True)
        p.add_run().add_picture(str(self.base / img["arquivo"]), width=Cm(float(img.get("largura_cm", 8))))

    def questao(self, numero, q):
        """Enunciado, figuras, alternativas/itens e espaço de resposta ficam juntos na mesma página."""
        paragrafos = paragrafos_de(q["enunciado"])
        valor = f" ({formatar_br(decimal_br(q['valor'], 'valor'))})"
        imagens = q.get("imagens") or ([q["imagem"]] if q.get("imagem") else [])
        for i, texto in enumerate(paragrafos):
            p = self.paragrafo(WD_ALIGN_PARAGRAPH.JUSTIFY, manter_com_proximo=True)
            if i == 0:
                adicionar_run(p, f"Questão {numero} – ", negrito=True)
                if q.get("rotulo"):
                    adicionar_run(p, f"{q['rotulo'].strip()} ")
            escrever_texto(p, texto, self.formulas)
            if i == len(paragrafos) - 1:
                adicionar_run(p, valor)
        for img in imagens:
            self.imagem(img)
        if q.get("alternativas"):
            lista = self.nova_lista()
            for alt in q["alternativas"]:
                p = self.paragrafo(lista=lista, manter_com_proximo=True)
                escrever_texto(p, str(alt), self.formulas)
        if q.get("itens"):
            lista = self.nova_lista()
            for item in q["itens"]:
                p = self.paragrafo(WD_ALIGN_PARAGRAPH.JUSTIFY, lista=lista, manter_com_proximo=True)
                escrever_texto(p, str(item["texto"]), self.formulas)
                if item.get("valor") is not None:
                    adicionar_run(p, f" ({formatar_br(decimal_br(item['valor'], 'valor do item'))})")
                self.linhas_em_branco(int(item.get("espaco_linhas", 4)), manter_com_proximo=True)
        padrao = 3 if q.get("tipo") == "objetiva" else (0 if q.get("itens") else 7)
        self.linhas_em_branco(max(int(q.get("espaco_linhas", padrao)), 1), manter_com_proximo=True)
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = False


# ---------------------------------------------------------------- validação da especificação

def validar_spec(spec, base):
    erros, avisos = [], []
    for campo in ("tipo", "bimestre", "serie", "disciplina", "professor", "total_pontos", "questoes"):
        if not spec.get(campo):
            erros.append(f"campo obrigatório ausente: {campo}")
    if erros:
        return erros, avisos, {}
    politica = spec.get("politica_questoes", "fonte")
    if politica not in {"fonte", "autoral_autorizado"}:
        erros.append("politica_questoes deve ser 'fonte' ou 'autoral_autorizado'")
    soma = Decimal(0)
    qtd_alternativas = set()
    distribuicao = {}
    tem_situacao_problema = False
    imagens = []
    for n, q in enumerate(spec["questoes"], 1):
        rot = f"Questão {n}"
        tipo = q.get("tipo")
        if tipo not in {"objetiva", "discursiva"}:
            erros.append(f"{rot}: tipo deve ser 'objetiva' ou 'discursiva'")
        if not paragrafos_de(q.get("enunciado")):
            erros.append(f"{rot}: enunciado vazio")
        try:
            valor = decimal_br(q.get("valor"), f"{rot} valor")
            if valor <= 0:
                raise ValueError(f"{rot}: valor deve ser positivo")
            soma += valor
            if q.get("itens") and all(i.get("valor") is not None for i in q["itens"]):
                soma_itens = sum(decimal_br(i["valor"], f"{rot} item") for i in q["itens"])
                if soma_itens != valor:
                    erros.append(f"{rot}: itens somam {formatar_br(soma_itens)}, valor da questão é {formatar_br(valor)}")
        except ValueError as e:
            erros.append(str(e))
        if q.get("autoral"):
            if politica != "autoral_autorizado":
                erros.append(f"{rot}: questão autoral sem autorização (politica_questoes: fonte)")
            else:
                avisos.append(f"{rot}: questão autoral — exige revisão do professor antes da aplicação")
        elif not str(q.get("fonte", "")).strip():
            erros.append(f"{rot}: sem 'fonte' (proveniência obrigatória: livro/capítulo/página, caderno ou banca/ano)")
        if tipo == "objetiva":
            alternativas = q.get("alternativas") or []
            if not 2 <= len(alternativas) <= len(LETRAS):
                erros.append(f"{rot}: objetiva precisa de 2 a {len(LETRAS)} alternativas")
            qtd_alternativas.add(len(alternativas))
            gabarito = str(q.get("gabarito", "")).strip().lower().rstrip(")")
            if len(gabarito) != 1 or gabarito not in LETRAS[:len(alternativas)]:
                erros.append(f"{rot}: gabarito ausente ou fora das alternativas")
            else:
                distribuicao[gabarito] = distribuicao.get(gabarito, 0) + 1
        elif tipo == "discursiva" and not (q.get("gabarito") or q.get("resposta_esperada")):
            avisos.append(f"{rot}: discursiva sem resposta esperada no gabarito do professor")
        if not q.get("resolucao"):
            avisos.append(f"{rot}: sem resolução registrada (a auditoria científica precisa da resolução)")
        texto_total = " ".join(paragrafos_de(q.get("enunciado")) + [str(a) for a in q.get("alternativas") or []])
        if re.search(r"formul[áa]rio|gloss[áa]rio", texto_total, re.I):
            erros.append(f"{rot}: o modelo proíbe fórmulas/glossário na prova")
        if q.get("situacao_problema"):
            tem_situacao_problema = True
        for img in q.get("imagens") or ([q["imagem"]] if q.get("imagem") else []):
            caminho = base / str(img.get("arquivo", ""))
            if not caminho.is_file():
                erros.append(f"{rot}: imagem não encontrada: {caminho}")
                continue
            largura = float(img.get("largura_cm", 8))
            px = DocxImage.from_file(str(caminho)).px_width
            dpi = px / (largura / 2.54)
            imagens.append({"questao": n, "arquivo": str(img["arquivo"]), "dpi_efetivo": round(dpi)})
            if dpi < 150:
                erros.append(f"{rot}: imagem {img['arquivo']} com {dpi:.0f} dpi no tamanho impresso (mínimo 150)")
            elif dpi < 220:
                avisos.append(f"{rot}: imagem {img['arquivo']} com {dpi:.0f} dpi (recomendado ≥ 220)")
            origem = str(img.get("origem", "")).lower()
            if not origem:
                avisos.append(f"{rot}: imagem sem 'origem' declarada")
            elif "foto" in origem:
                erros.append(f"{rot}: o modelo não aceita fotos de documentos/livros")
    total = decimal_br(spec["total_pontos"], "total_pontos")
    if soma != total:
        erros.append(f"soma dos valores ({formatar_br(soma)}) diferente do total de pontos ({formatar_br(total)})")
    if len(qtd_alternativas) > 1:
        avisos.append(f"objetivas com quantidades diferentes de alternativas: {sorted(qtd_alternativas)}")
    segmento = str(spec.get("segmento", "medio")).lower()
    if segmento in SEGMENTOS_COM_SITUACAO_PROBLEMA and not tem_situacao_problema:
        erros.append("o modelo exige ao menos uma questão com situação-problema (marque situacao_problema: true)")
    if distribuicao:
        maior = max(distribuicao.values())
        if len(distribuicao) > 1 and maior > 2 * min(distribuicao.values()) + 1:
            avisos.append(f"gabarito concentrado: {dict(sorted(distribuicao.items()))}")
    return erros, avisos, {"soma": formatar_br(soma), "distribuicao_gabarito": dict(sorted(distribuicao.items())),
                           "imagens": imagens}


# ---------------------------------------------------------------- saída

def docx_sem_fontes_embutidas(origem, destino):
    """Remove as fontes TrueType embutidas (~3 MB). Tahoma/Arial existem no Windows, macOS e Google Docs."""
    with zipfile.ZipFile(origem) as zin, zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename.startswith("word/fonts/"):
                continue
            dados = zin.read(item.filename)
            if item.filename == "word/fontTable.xml":
                dados = re.sub(rb"<w:embed(Regular|Bold|Italic|BoldItalic)\b[^>]*/>", b"", dados)
            elif item.filename == "word/_rels/fontTable.xml.rels":
                dados = re.sub(rb"<Relationship\b[^>]*/>", b"", dados)
            elif item.filename == "word/settings.xml":
                dados = dados.replace(b"<w:embedTrueTypeFonts/>", b"")
            zout.writestr(item, dados)


def libreoffice_math_instalado():
    """Sem o módulo Math, o LibreOffice descarta as equações OMML ao gerar o PDF."""
    soffice = shutil.which("soffice")
    if not soffice:
        return False
    raiz = Path(soffice).resolve().parent.parent
    return any(raiz.rglob("libsmlo.*"))


def renderizar_pdf(docx, pasta):
    perfil = Path(tempfile.mkdtemp(prefix="lo-mater-"))
    try:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{perfil}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(pasta), str(docx)],
                       check=True, capture_output=True, timeout=180)
    finally:
        shutil.rmtree(perfil, ignore_errors=True)
    return Path(pasta) / (Path(docx).stem + ".pdf")


def auditar_pdf(pdf, spec, erros, avisos):
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    paginas = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    fontes = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    if "Tahoma" not in fontes:
        avisos.append("o PDF foi renderizado sem Tahoma (fonte substituta); confira o DOCX no Word")
    texto = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
    for campo in ("XºAno", "Xª Série", "MATÉRIA", "SEU NOME", "TESTE/PROVA", "3,0/7,0", "/03/2026"):
        if campo in texto:
            erros.append(f"campo do modelo não preenchido no PDF: {campo!r}")
    posicoes = []
    for n in range(1, len(spec["questoes"]) + 1):
        achados = [m.start() for m in re.finditer(rf"Questão {n} –", texto)]
        if len(achados) != 1:
            erros.append(f"'Questão {n} –' aparece {len(achados)} vez(es) no PDF")
        else:
            posicoes.append(achados[0])
    if posicoes != sorted(posicoes):
        erros.append("questões fora de ordem no PDF")
    return paginas


def gerar_previas(pdf, pasta):
    pasta.mkdir(parents=True, exist_ok=True)
    for antigo in pasta.glob("pagina-*.png"):
        antigo.unlink()
    subprocess.run(["pdftoppm", "-r", "80", "-png", str(pdf), str(pasta / "pagina")], check=True)
    return sorted(pasta.glob("pagina-*.png"))


def gerar_gabarito(spec, formulas, destino):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(int(PAGINA["w"])), Twips(int(PAGINA["h"]))
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(2))
    estilo = doc.styles["Normal"]
    estilo.font.name = FONTE_CORPO
    estilo.font.size = Pt(11)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE_CORPO)
    tipo = str(spec.get("tipo", "TESTE")).upper()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GABARITO DO PROFESSOR — USO INTERNO, NÃO DISTRIBUIR")
    r.bold = True
    doc.add_paragraph(f"Escola Mater / Colégio São Conrado — {tipo} {spec['bimestre']}º BIM — {spec['serie']} — "
                      f"{spec['disciplina'].upper()} — Prof. {spec['professor']} — Total: "
                      f"{formatar_br(decimal_br(spec['total_pontos'], 'total_pontos'))}")
    tabela = doc.add_table(rows=1, cols=4)
    tabela.style = "Table Grid"
    for celula, texto in zip(tabela.rows[0].cells, ("Questão", "Valor", "Resposta", "Fonte")):
        celula.paragraphs[0].add_run(texto).bold = True
    for n, q in enumerate(spec["questoes"], 1):
        resposta = str(q.get("gabarito") or q.get("resposta_esperada") or "—")
        if q.get("tipo") == "objetiva":
            resposta = resposta.strip().lower().rstrip(")") + ")"
        linha = tabela.add_row().cells
        linha[0].text = str(n)
        linha[1].text = formatar_br(decimal_br(q["valor"], "valor"))
        escrever_texto(linha[2].paragraphs[0], resposta, formulas)
        linha[3].text = "AUTORAL (revisar)" if q.get("autoral") else str(q.get("fonte", ""))
    tabela.autofit = False
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, (1.8, 1.6, 7.6, 6.0)):
            celula.width = Cm(largura)
    doc.add_paragraph()
    titulo = doc.add_paragraph()
    titulo.add_run("Resoluções").bold = True
    for n, q in enumerate(spec["questoes"], 1):
        if not q.get("resolucao"):
            continue
        for i, texto in enumerate(paragrafos_de(q["resolucao"])):
            p = doc.add_paragraph()
            if i == 0:
                adicionar_run(p, f"Questão {n} – ", negrito=True)
            escrever_texto(p, texto, formulas)
    doc.save(destino)


RENDER_SVG_JS = """
const { chromium } = require('playwright');
const fs = require('fs');
(async () => {
  const [svg, png, w, h, escala] = process.argv.slice(1);
  const browser = await chromium.launch();
  const page = await browser.newPage({ viewport: { width: +w, height: +h }, deviceScaleFactor: +escala });
  await page.setContent('<html><body style="margin:0;background:#fff">' + fs.readFileSync(svg, 'utf8') + '</body></html>');
  await page.screenshot({ path: png, clip: { x: 0, y: 0, width: +w, height: +h } });
  await browser.close();
})().catch(e => { console.error(e.message); process.exit(1); });
"""


def comando_figura(args):
    from PIL import Image

    svg = Path(args.svg)
    tag = re.search(r"<svg\b[^>]*>", svg.read_text(encoding="utf-8"))
    medidas = {k: re.search(rf'\b{k}="([\d.]+)', tag.group(0)) if tag else None for k in ("width", "height")}
    if not all(medidas.values()):
        raise ErroModelo("o <svg> precisa de width e height em px")
    largura, altura = (float(medidas[k].group(1)) for k in ("width", "height"))
    escala = args.largura_cm / 2.54 * args.dpi / largura
    png = svg.with_suffix(".png")
    try:
        subprocess.run(["node", "-e", RENDER_SVG_JS, str(svg), str(png), str(largura), str(altura), f"{escala:.4f}"],
                       check=True, capture_output=True, text=True, timeout=120)
    except (FileNotFoundError, subprocess.CalledProcessError) as e:
        raise ErroModelo(f"falha ao renderizar com Node + Playwright (npm i -g playwright): {getattr(e, 'stderr', e)}")
    imagem = Image.open(png)
    imagem = imagem.convert("RGB").quantize(64) if args.cores else imagem.convert("L").quantize(16, dither=Image.Dither.NONE)
    imagem.save(png, optimize=True, dpi=(args.dpi, args.dpi))
    largura_px = imagem.size[0]
    print(json.dumps({"png": str(png), "px": list(imagem.size), "bytes": png.stat().st_size,
                      "dpi_efetivo": round(largura_px / (args.largura_cm / 2.54)),
                      "fonte_minima_px_para_9pt": round(0.3175 * largura / args.largura_cm)},
                     ensure_ascii=False, indent=2))
    return 0


def comando_modelo(args):
    destino = Path(args.saida)
    decodificar_modelo(args.origem, destino)
    erros, avisos = conferir_modelo(destino)
    print(json.dumps({"modelo": str(destino), "sha256": sha256(destino), "erros": erros, "avisos": avisos},
                     ensure_ascii=False, indent=2))
    return 1 if erros else 0


def comando_gerar(args):
    spec_path = Path(args.spec)
    spec = carregar_spec(spec_path)
    base = spec_path.parent
    if not Path(args.modelo).is_file():
        raise ErroModelo(f"modelo não encontrado em {args.modelo}: baixe-o do Drive (file ID {DRIVE_FILE_ID}) "
                         "e rode o subcomando 'modelo'")
    erros, avisos, resumo = validar_spec(spec, base)
    erros_modelo, avisos_modelo = conferir_modelo(args.modelo)
    erros += erros_modelo
    avisos += [f"modelo: {a}" for a in avisos_modelo]
    if erros or args.so_validar:
        print(json.dumps({"gate": "BLOQUEADO" if erros else "ESPECIFICAÇÃO OK", "erros": erros,
                          "avisos": avisos, **resumo}, ensure_ascii=False, indent=2))
        return 1 if erros else 0

    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)
    nome = nome_base(spec)
    textos = []
    for q in spec["questoes"]:
        textos += paragrafos_de(q.get("enunciado")) + [str(a) for a in q.get("alternativas") or []]
        textos += [str(i.get("texto", "")) for i in q.get("itens") or []]
        textos += paragrafos_de(q.get("resolucao")) + [str(q.get("gabarito") or q.get("resposta_esperada") or "")]
    formulas = converter_formulas([f for t in textos for f in formulas_do_texto(t)])
    if formulas and not libreoffice_math_instalado():
        raise ErroModelo("há fórmulas \\( \\), mas o LibreOffice está sem o módulo Math e as apagaria do PDF: "
                         "sudo apt-get install -y libreoffice-math")

    doc = Document(args.modelo)
    preencher_cabecalho(doc, spec, preto=args.cabecalho_preto)
    cabecalho_so_na_primeira_pagina(doc)
    remover_exemplo_do_corpo(doc)
    corpo = Corpo(doc, formulas, base)
    corpo.linhas_em_branco(1)
    for n, q in enumerate(spec["questoes"], 1):
        corpo.questao(n, q)

    with tempfile.TemporaryDirectory() as tmp:
        com_fontes = Path(tmp) / f"{nome}.docx"
        doc.save(com_fontes)
        pdf = renderizar_pdf(com_fontes, tmp)  # renderiza com as fontes do modelo embutidas
        shutil.move(str(pdf), saida / f"{nome}.pdf")
        docx_final = saida / f"{nome}.docx"
        if args.docx_com_fontes:
            shutil.copy(com_fontes, docx_final)
        else:
            docx_sem_fontes_embutidas(com_fontes, docx_final)
    pdf = saida / f"{nome}.pdf"
    paginas = auditar_pdf(pdf, spec, erros, avisos)
    previas = gerar_previas(pdf, saida / "previa")

    gab_docx = saida / f"{nome} – GABARITO PROFESSOR.docx"
    gerar_gabarito(spec, formulas, gab_docx)
    gab_pdf = renderizar_pdf(gab_docx, saida)

    manifesto = {
        "gate": "BLOQUEADO" if erros else "APROVADO (falta inspeção visual de todas as páginas)",
        "gerado_em": dt.datetime.now().isoformat(timespec="seconds"),
        "modelo": {"drive_file_id": DRIVE_FILE_ID, "titulo": DRIVE_TITULO, "sha256": sha256(args.modelo)},
        "avaliacao": {k: spec.get(k) for k in ("tipo", "bimestre", "serie", "disciplina", "professor",
                                               "data", "total_pontos", "segmento", "politica_questoes")},
        "paginas": paginas,
        **resumo,
        "questoes": [{"numero": n, "tipo": q.get("tipo"), "valor": q.get("valor"), "fonte": q.get("fonte"),
                      "rotulo": q.get("rotulo"), "autoral": bool(q.get("autoral")),
                      "situacao_problema": bool(q.get("situacao_problema")),
                      "gabarito": q.get("gabarito") or q.get("resposta_esperada")}
                     for n, q in enumerate(spec["questoes"], 1)],
        "arquivos": {p.name: {"sha256": sha256(p), "bytes": p.stat().st_size}
                     for p in (docx_final, pdf, gab_docx, gab_pdf)},
        "previas": [str(p) for p in previas],
        "erros": erros,
        "avisos": avisos,
    }
    caminho_manifesto = saida / f"{nome} – manifesto.json"
    caminho_manifesto.write_text(json.dumps(manifesto, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"gate": manifesto["gate"], "paginas": paginas, "saida": str(saida),
                      "arquivos": list(manifesto["arquivos"]), "previas": manifesto["previas"],
                      "erros": erros, "avisos": avisos}, ensure_ascii=False, indent=2))
    return 1 if erros else 0


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="comando", required=True)
    m = sub.add_parser("modelo", help="decodifica e confere o modelo oficial baixado do Drive")
    m.add_argument("origem", help=".docx ou JSON salvo pelo download_file_content do MCP do Drive")
    m.add_argument("--saida", required=True, help="caminho do .docx do modelo a gravar")
    m.set_defaults(func=comando_modelo)
    g = sub.add_parser("gerar", help="gera a avaliação a partir de uma especificação JSON/YAML")
    g.add_argument("spec")
    g.add_argument("--modelo", required=True, help="MODELO AVALIAÇÕES - Fund II e Médio.docx")
    g.add_argument("--saida", required=True, help="pasta de saída")
    g.add_argument("--so-validar", action="store_true", help="só valida a especificação e o modelo")
    g.add_argument("--cabecalho-preto", action="store_true", help="campos preenchidos em preto (padrão: cor do modelo)")
    g.add_argument("--docx-com-fontes", action="store_true", help="mantém as fontes embutidas do modelo (~3 MB)")
    g.set_defaults(func=comando_gerar)
    f = sub.add_parser("figura", help="renderiza uma ilustração SVG em PNG de alta resolução")
    f.add_argument("svg", help="arquivo .svg com width/height em px; o PNG sai ao lado")
    f.add_argument("--largura-cm", type=float, required=True, help="largura que a figura terá na prova")
    f.add_argument("--dpi", type=int, default=400)
    f.add_argument("--cores", action="store_true", help="mantém cores (padrão: 16 tons de cinza)")
    f.set_defaults(func=comando_figura)
    args = parser.parse_args()
    try:
        sys.exit(args.func(args))
    except (ErroModelo, ValueError, OSError, subprocess.CalledProcessError) as e:
        print(json.dumps({"gate": "BLOQUEADO", "erros": [str(e)]}, ensure_ascii=False, indent=2))
        sys.exit(1)


if __name__ == "__main__":
    main()
